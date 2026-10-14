import openpyxl
from app.core.database import SessionLocal
from app.models.skills import Skill
from sqlalchemy import select
from pathlib import Path

DATA_FILE = Path(__file__).resolve().parent.parent / "data" / "BB - Database.xlsx"

def main():
    wb = openpyxl.load_workbook(DATA_FILE, data_only=True)
    ws = wb["Skills"]
    
    headers = [cell.value for cell in ws[1]]

    added = 0
    with SessionLocal() as db:
        for col_idx, header in enumerate(headers):
            if col_idx == 0 or header is None:
                continue

            category = header.strip().title() 
            if category.startswith("Elite"):
                continue
            for (value,) in ws.iter_rows(
                min_row=2, min_col=col_idx + 1, max_col=col_idx + 1, values_only=True
            ):
                if value is None or not str(value).strip():
                    continue 
                name = str(value).strip()

                existing = db.scalar(select(Skill).where(Skill.name == name))
                if existing:
                    existing.category = category 
                else:
                    db.add(Skill(name=name, category=category))
                    added += 1
        db.commit()
        
    print(f"Done. {added} skills added.")

if __name__ == "__main__":
    main()